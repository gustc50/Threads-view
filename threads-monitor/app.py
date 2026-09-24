"""
Interface web local para o monitor do Threads.

Rode com:
    python app.py

Abre um servidor em http://127.0.0.1:5000 com tres abas:
    - Configuracoes: token, app id/secret, tipo, max de paginas, palavras-chave
    - Execucao: botao "Rodar agora" + log ao vivo
    - Resultados: tabela com os links coletados (lida de resultados.xlsx)

Roda apenas localmente (127.0.0.1). O token nunca sai da sua maquina.
"""

from __future__ import annotations

import threading
import webbrowser
from datetime import datetime

from flask import Flask, jsonify, render_template, request
from openpyxl import load_workbook

from collector import run_collection
from diagnostico import rodar_diagnostico
from settings import (
    OUTPUT_XLSX,
    load_config,
    load_keywords,
    save_env,
    save_keywords,
)
from storage import COLUMNS

app = Flask(__name__)

# Estado da execucao em andamento (protegido por lock)
_estado = {
    "rodando": False,
    "logs": [],          # lista de {"hora": str, "linha": str}
    "resumo": None,      # dict devolvido por run_collection
    "iniciado_em": None,
}
_lock = threading.Lock()


def _log(linha: str) -> None:
    with _lock:
        _estado["logs"].append({
            "hora": datetime.now().strftime("%H:%M:%S"),
            "linha": linha,
        })


def _executar_coleta(search_type: str) -> None:
    """Roda a coleta numa thread separada, alimentando o log."""
    try:
        cfg = load_config()
        if not cfg.is_valid():
            _log("ERRO: token nao configurado. Salve o token na aba Configuracoes.")
            return
        keywords = load_keywords()
        if not keywords:
            _log("ERRO: nenhuma palavra-chave. Adicione ao menos uma na aba Configuracoes.")
            return
        resumo = run_collection(
            cfg=cfg,
            keywords=keywords,
            search_type=search_type,
            log_fn=_log,
        )
        with _lock:
            _estado["resumo"] = resumo
    except Exception as exc:  # noqa: BLE001 - queremos reportar tudo no log
        _log(f"ERRO inesperado: {exc}")
    finally:
        with _lock:
            _estado["rodando"] = False


@app.route("/")
def index():
    return render_template("index.html")


@app.route("/api/config", methods=["GET"])
def get_config():
    cfg = load_config()
    return jsonify({
        # Nao devolvemos o token inteiro; so indicamos se existe
        "tem_token": bool(cfg.access_token),
        "token_mascarado": _mascara(cfg.access_token),
        "app_id": cfg.app_id,
        "tem_app_secret": bool(cfg.app_secret),
        "max_pages": cfg.max_pages,
        "search_type": cfg.search_type,
        "keywords": load_keywords(),
    })


@app.route("/api/config", methods=["POST"])
def post_config():
    dados = request.get_json(force=True) or {}
    env_updates: dict[str, str] = {}

    # So sobrescreve o token/secret se o usuario digitou algo novo
    if dados.get("access_token"):
        env_updates["THREADS_ACCESS_TOKEN"] = dados["access_token"].strip()
    if dados.get("app_id") is not None:
        env_updates["THREADS_APP_ID"] = str(dados["app_id"]).strip()
    if dados.get("app_secret"):
        env_updates["THREADS_APP_SECRET"] = dados["app_secret"].strip()
    if dados.get("max_pages") is not None:
        env_updates["MAX_PAGES"] = str(dados["max_pages"]).strip()
    if dados.get("search_type"):
        st = str(dados["search_type"]).upper()
        env_updates["SEARCH_TYPE"] = st if st in ("RECENT", "TOP") else "RECENT"

    if env_updates:
        save_env(env_updates)

    if "keywords" in dados and isinstance(dados["keywords"], list):
        save_keywords(dados["keywords"])

    return jsonify({"ok": True})


@app.route("/api/diagnostico", methods=["POST"])
def diagnostico():
    cfg = load_config()
    etapas = rodar_diagnostico(cfg.access_token)
    return jsonify({"etapas": etapas})


@app.route("/api/run", methods=["POST"])
def run():
    with _lock:
        if _estado["rodando"]:
            return jsonify({"ok": False, "erro": "Ja existe uma coleta em andamento."}), 409
        _estado["rodando"] = True
        _estado["logs"] = []
        _estado["resumo"] = None
        _estado["iniciado_em"] = datetime.now().strftime("%H:%M:%S")

    dados = request.get_json(force=True) or {}
    search_type = str(dados.get("search_type", "")).upper()
    if search_type not in ("RECENT", "TOP"):
        search_type = load_config().search_type

    t = threading.Thread(target=_executar_coleta, args=(search_type,), daemon=True)
    t.start()
    return jsonify({"ok": True})


@app.route("/api/status")
def status():
    with _lock:
        return jsonify({
            "rodando": _estado["rodando"],
            "logs": _estado["logs"],
            "resumo": _estado["resumo"],
            "iniciado_em": _estado["iniciado_em"],
        })


@app.route("/api/results")
def results():
    """Le resultados.xlsx e devolve as linhas (mais recentes primeiro)."""
    if not OUTPUT_XLSX.exists():
        return jsonify({"colunas": COLUMNS, "linhas": []})

    wb = load_workbook(OUTPUT_XLSX, read_only=True)
    ws = wb.active
    linhas = []
    for i, row in enumerate(ws.iter_rows(values_only=True)):
        if i == 0:
            continue  # cabecalho
        linhas.append([("" if c is None else str(c)) for c in row])
    wb.close()

    linhas.reverse()  # mais recentes primeiro
    return jsonify({"colunas": COLUMNS, "linhas": linhas})


def _mascara(token: str) -> str:
    if not token:
        return ""
    if len(token) <= 8:
        return "*" * len(token)
    return f"{token[:4]}...{token[-4:]}"


def main() -> None:
    url = "http://127.0.0.1:5000"
    # Abre o navegador automaticamente (levemente atrasado para o servidor subir)
    threading.Timer(1.0, lambda: webbrowser.open(url)).start()
    print(f"Threads Monitor rodando em {url} (Ctrl+C para encerrar)")
    app.run(host="127.0.0.1", port=5000, debug=False)


if __name__ == "__main__":
    main()
