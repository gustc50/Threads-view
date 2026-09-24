"""
Modulo responsavel por:
    - carregar/salvar ids ja vistos (seen_ids.json);
    - gravar/atualizar a planilha resultados.xlsx.
"""

from __future__ import annotations

import json
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook, load_workbook

COLUMNS = [
    "data_coleta",
    "palavra_chave",
    "usuario",
    "data_publicacao",
    "texto",
    "link",
]

TEXT_LIMIT = 200


def load_seen_ids(path: Path) -> set[str]:
    """Le o arquivo seen_ids.json e devolve um set de ids."""
    if not path.exists():
        return set()
    try:
        with path.open("r", encoding="utf-8") as fh:
            data = json.load(fh)
        if isinstance(data, list):
            return {str(x) for x in data}
        return set()
    except (json.JSONDecodeError, OSError):
        return set()


def save_seen_ids(path: Path, ids: set[str]) -> None:
    """Salva os ids no disco."""
    with path.open("w", encoding="utf-8") as fh:
        json.dump(sorted(ids), fh, ensure_ascii=False, indent=2)


def _ensure_workbook(path: Path) -> Workbook:
    """Abre a planilha existente ou cria uma nova com cabecalho."""
    if path.exists():
        return load_workbook(path)

    wb = Workbook()
    ws = wb.active
    ws.title = "publicacoes"
    ws.append(COLUMNS)
    return wb


def append_rows(path: Path, rows: list[dict]) -> None:
    """Anexa linhas na planilha. Cria o arquivo se nao existir."""
    if not rows:
        return

    wb = _ensure_workbook(path)
    ws = wb.active

    for row in rows:
        ws.append([row.get(col, "") for col in COLUMNS])

    wb.save(path)


def build_row(keyword: str, post: dict) -> dict:
    """Converte um post da API em uma linha da planilha."""
    text = (post.get("text") or "").replace("\n", " ").strip()
    if len(text) > TEXT_LIMIT:
        text = text[:TEXT_LIMIT].rstrip() + "..."

    return {
        "data_coleta": datetime.utcnow().strftime("%Y-%m-%d %H:%M:%S"),
        "palavra_chave": keyword,
        "usuario": post.get("username", ""),
        "data_publicacao": post.get("timestamp", ""),
        "texto": text,
        "link": post.get("permalink", ""),
    }


def filter_new_posts(
    posts: list[dict],
    seen: set[str],
) -> tuple[list[dict], set[str]]:
    """
    Retorna (posts_novos, ids_novos).
    O caller decide quando persistir seen_ids.
    """
    novos: list[dict] = []
    novos_ids: set[str] = set()

    for post in posts:
        pid = str(post.get("id") or "")
        if not pid or pid in seen or pid in novos_ids:
            continue
        novos.append(post)
        novos_ids.add(pid)

    return novos, novos_ids
